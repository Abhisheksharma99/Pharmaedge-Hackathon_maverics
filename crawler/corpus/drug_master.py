"""
Drug master: the team's canonical drug list (AdisInsight-derived spreadsheet, ~75k drugs) as a lookup corpus,
`pharmaedge.drug_master` - one doc per drug: AdisInsight id, name, every alternative name (brands, code names),
companies, mechanism, targets, modality, therapy areas.

Used for RESOLUTION only: which drug a name means ("Winrevair", "MK-7962", "ACE-011" -> sotatercept, Adis
800024655), across the platform's asset slugs and patent_intel's Adis ids. Master names are never added to an asset's
crawl aliases: the crawler searches every source for each alias, and a master lists dozens of internal codes.

Load (once per spreadsheet version; idempotent):  cd crawler && python -m scripts.load_drug_master <file.xlsx>

Spreadsheet quirks handled here: PostgreSQL array literals ("{a,b,\"c, d\"}"), and names qualified by company
("Treprostinil - United Therapeutics Corporation", "Sotatercept-csrk - Merck & Co"): the part before " - " is the
name. Lookup keys are the names lower-cased with everything but letters and digits removed ("BI 1015550",
"BI-1015550" and "bi1015550" are one key); keys shorter than 3 characters are not indexed.
"""

import os
import re
from datetime import datetime, timezone
from typing import Any, Dict, Iterable, Iterator, List, Optional

from pymongo import UpdateOne

from corpus import collection

COLLECTION = os.getenv("DRUG_MASTER_CORPUS", "pharmaedge.drug_master")
ADIS = re.compile(r"\d{9}")
MAX_NAMES = 80


def pg_array(value: Any) -> List[str]:
    """'{a,b,"c, d"}' -> ['a', 'b', 'c, d']; a plain string -> [string]; empty -> []."""
    if value is None:
        return []
    s = str(value).strip()
    if not s:
        return []
    if not (s.startswith("{") and s.endswith("}")):
        return [s]
    parts = re.findall(r'"((?:[^"\\]|\\.)*)"|([^,]+)', s[1:-1])
    out = [(quoted or bare).strip() for quoted, bare in parts]
    return [x for x in out if x and x.upper() != "NULL"]


def base_name(name: str) -> str:
    """'Treprostinil - United Therapeutics Corporation' -> 'Treprostinil' (company qualifier dropped)."""
    return re.split(r"\s+-\s+", name.strip(), maxsplit=1)[0].strip()


def key(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", name.lower())


def to_doc(row: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    name = str(row.get("drug_name") or "").strip()
    if not name:
        return None
    adis = str(row.get("drug_web_id") or "").strip()
    adis = adis if ADIS.fullmatch(adis) else None
    names = list(dict.fromkeys(n for n in (base_name(x) for x in [name, str(row.get("competitor_name") or ""),
                                                                 *pg_array(row.get("alternative_drug_name"))]) if n))[:MAX_NAMES]
    return {
        "_id": adis or f"row:{row.get('drug_id')}",
        "adis_id": adis,
        "name": base_name(name),
        "full_name": name,
        "names": names,
        "keys": sorted({k for k in map(key, names) if len(k) >= 3}),
        "companies": pg_array(row.get("competitor_company")),
        "therapy_areas": pg_array(row.get("competitor_primary_therapy_area")),
        "moa": pg_array(row.get("competitor_moa")),
        "targets": pg_array(row.get("competitor_target")),
        "modality": pg_array(row.get("competitor_modality")),
        "drug_type": pg_array(row.get("competitor_drug_type")),
        "source": row.get("source"),
    }


def rows(path: str) -> Iterator[Dict[str, Any]]:
    """Rows of the 'Drug Master' sheet as dicts (streamed: the sheet has ~75k rows)."""
    from openpyxl import load_workbook  # only the loader needs it

    ws = load_workbook(path, read_only=True, data_only=True)["Drug Master"]
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    for r in it:
        if r and r[0] is not None:
            yield dict(zip(header, r))


def store(db, docs: Iterable[Dict[str, Any]], batch: int = 1000) -> Dict[str, int]:
    coll = collection(db, COLLECTION)
    coll.create_index("keys")
    coll.create_index("adis_id")
    now, counts, ops = datetime.now(timezone.utc), {"drugs": 0, "new": 0}, []

    def flush():
        if ops:
            counts["new"] += coll.bulk_write(ops, ordered=False).upserted_count
            ops.clear()

    for d in docs:
        counts["drugs"] += 1
        ops.append(UpdateOne({"_id": d["_id"]}, {"$set": {**d, "loaded_at": now}}, upsert=True))
        if len(ops) >= batch:
            flush()
    flush()
    return counts


def lookup(db, name: str) -> List[Dict[str, Any]]:
    """Master drugs carrying this name (exact, by key). Several when companies each have their own entry."""
    k = key(base_name(name))
    return list(collection(db, COLLECTION).find({"keys": k})) if len(k) >= 3 else []


GENERIC_COMPANY_WORDS = {"pharma", "pharmaceutical", "pharmaceuticals", "therapeutics", "biosciences", "biopharma", "biotech",
                         "corporation", "company", "group", "holdings", "international", "laboratories", "sciences", "medicines"}


def _company_words(name: str) -> set:
    return {w for w in re.findall(r"[a-z0-9]+", name.lower()) if len(w) >= 4 and w not in GENERIC_COMPANY_WORDS}


def resolve(db, names: List[str], company: Optional[str] = None) -> Optional[Dict[str, Any]]:
    """The one master drug these names mean (names[0] is the canonical name), or None when it cannot be told.

    A substance usually has one entry per developing company and formulation. With a single candidate, that is it.
    Otherwise only candidates of the asset's company count (a distinctive company word in common: "Actelion
    (Janssen)" ~ "Actelion Pharmaceuticals Ltd"), ranked by: master name equal to the canonical name
    ("Treprostinil", not "Treprostinil dry powder inhalation"), then how many of the asset's names it carries.
    A tie at the top is not guessed."""
    found: Dict[str, Dict[str, Any]] = {}
    hits: Dict[str, int] = {}
    for n in names:
        for d in lookup(db, n):
            found[d["_id"]] = d
            hits[d["_id"]] = hits.get(d["_id"], 0) + 1
    if len(found) == 1:
        return next(iter(found.values()))
    words = _company_words(company or "")
    canonical = key(base_name(names[0])) if names else ""
    ranked = sorted(((key(d["name"]) == canonical, hits[i]), i) for i, d in found.items()
                    if words and any(words & _company_words(c) for c in d.get("companies", [])))
    if not ranked or (len(ranked) > 1 and ranked[-1][0] == ranked[-2][0]):
        return None
    return found[ranked[-1][1]]
